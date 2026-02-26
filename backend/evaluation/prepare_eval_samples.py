#!/usr/bin/env python3
import argparse
import shutil
import time
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def mutate_text_file(path: Path, token: str) -> None:
    text = path.read_text(encoding='utf-8', errors='ignore')
    text += f"\n\nrun token: {token}\n"
    path.write_text(text, encoding='utf-8')


def mutate_docx(path: Path, token: str) -> None:
    doc = Document(str(path))
    doc.add_paragraph(f"run token: {token}")
    doc.save(str(path))


def mutate_xlsx(path: Path, token: str) -> None:
    wb = load_workbook(str(path))
    ws = wb.active
    ws.append(["run_token", token, "e2e_eval"])
    wb.save(str(path))


def run() -> int:
    ap = argparse.ArgumentParser(description='Prepare unique ingestion samples for repeatable regression.')
    ap.add_argument('--source', default='/app/ingest_samples')
    ap.add_argument('--target', default='/app/ingest_eval/current')
    ap.add_argument('--run-id', default='')
    args = ap.parse_args()

    source = Path(args.source)
    target = Path(args.target)
    token = args.run_id.strip() or time.strftime('%Y%m%d%H%M%S')

    if not source.exists() or not source.is_dir():
        raise SystemExit(f'source_not_found:{source}')

    if target.exists():
        shutil.rmtree(target)
    target.mkdir(parents=True, exist_ok=True)

    copied = 0
    for src in sorted(source.iterdir()):
        if not src.is_file():
            continue
        dst_name = f'run-{token}-{src.name}'
        dst = target / dst_name
        shutil.copy2(src, dst)

        ext = dst.suffix.lower()
        if ext in {'.txt', '.md'}:
            mutate_text_file(dst, token)
        elif ext == '.docx':
            mutate_docx(dst, token)
        elif ext == '.xlsx':
            mutate_xlsx(dst, token)
        copied += 1

    print('prepared_samples', copied)
    print('target', str(target))
    print('token', token)
    return 0


if __name__ == '__main__':
    raise SystemExit(run())
